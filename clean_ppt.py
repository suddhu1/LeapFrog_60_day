import os

# Automatically install python-pptx if you don't have it
try:
    from pptx import Presentation
except ImportError:
    print("Installing required library...")
    os.system('pip install python-pptx')
    from pptx import Presentation

def remove_gamma_watermark(input_file, output_file):
    if not os.path.exists(input_file):
        print(f"Error: Could not find file '{input_file}' in this folder.")
        return

    prs = Presentation(input_file)
    target_text = "made with gamma"
    removed_count = 0

    # 1. Clean individual slides
    for i, slide in enumerate(prs.slides):
        to_delete = []
        for shape in slide.shapes:
            if shape.has_text_frame and target_text in shape.text.lower():
                to_delete.append(shape)
        for shape in to_delete:
            sp = shape._element
            sp.getparent().remove(sp)
            removed_count += 1
            print(f"Removed watermark from Slide {i+1}")

    # 2. Clean slide masters and layouts (where watermarks are often hidden)
    for master in prs.slide_masters:
        to_delete = []
        for shape in master.shapes:
            if shape.has_text_frame and target_text in shape.text.lower():
                to_delete.append(shape)
        for shape in to_delete:
            sp = shape._element
            sp.getparent().remove(sp)
            removed_count += 1

        for layout in master.slide_layouts:
            to_delete = []
            for shape in layout.shapes:
                if shape.has_text_frame and target_text in shape.text.lower():
                    to_delete.append(shape)
            for shape in to_delete:
                sp = shape._element
                sp.getparent().remove(sp)
                removed_count += 1

    prs.save(output_file)
    print(f"\nSuccess! Removed {removed_count} instances of the watermark.")
    print(f"Cleaned presentation saved as: {output_file}")

# Adjust the filenames if needed
remove_gamma_watermark("Introduction-to-Vesicular-Transport .ppt", "Introduction-to-Vesicular-Transport_Cleaned.pptx")